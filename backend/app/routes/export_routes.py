from flask import Blueprint, jsonify, send_file, request
from app.database import get_connection
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

export_bp = Blueprint('export', __name__, url_prefix='/export')

# ============ FUNCIÓN AUXILIAR ============
def obtener_datos_arboles():
    """Obtiene todos los árboles con sus especies para exportar"""
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("""
        SELECT 
            a.MedicionArbolID AS ID,
            a.NumeroArbol AS Numero_Arbol,
            e.NombreComun AS Especie_Comun,
            e.NombreCientifico AS Especie_Cientifica,
            e.Familia,
            e.Categoria,
            a.DAP_M AS DAP_m,
            a.AlturaTotal_Mts AS Altura_Total_m,
            a.AlturaComercial_Mts AS Altura_Comercial_m,
            a.AreaBasal_M2 AS Area_Basal_m2,
            a.VolumenTotal_M3 AS Volumen_Total_m3,
            a.BiomasaAerea_kg AS Biomasa_kg,
            a.Carbono_kg AS Carbono_kg,
            a.CO2e_kg AS CO2e_kg,
            a.DensidadMadera AS Densidad_Madera,
            a.EstadoSanitario AS Estado_Sanitario,
            a.Latitud,
            a.Longitud,
            a.PuntoGPS AS Punto_GPS,
            a.Transecto,
            a.Observaciones
        FROM Info_arboles a
        INNER JOIN Especies e ON a.EspecieID = e.EspecieID
        ORDER BY a.MedicionArbolID
    """)
    datos = cursor.fetchall()
    cursor.close()
    conn.close()
    return datos

# ============ EXPORTAR ÁRBOLES ============
@export_bp.route('/arboles/csv', methods=['GET'])
def exportar_arboles_csv():
    datos = obtener_datos_arboles()
    
    if not datos:
        return jsonify({'error': 'No hay datos para exportar'}), 404
    
    # Crear CSV en memoria
    output = io.StringIO()
    # Escribir BOM para que Excel reconozca los acentos
    output.write('\ufeff')
    
    writer = csv.DictWriter(output, fieldnames=datos[0].keys())
    writer.writeheader()
    writer.writerows(datos)
    
    # Convertir a bytes
    output.seek(0)
    bytes_output = io.BytesIO(output.getvalue().encode('utf-8'))
    
    filename = f"arboles_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    
    return send_file(
        bytes_output,
        mimetype='text/csv',
        as_attachment=True,
        download_name=filename
    )

@export_bp.route('/arboles/xlsx', methods=['GET'])
def exportar_arboles_xlsx():
    datos = obtener_datos_arboles()
    
    if not datos:
        return jsonify({'error': 'No hay datos para exportar'}), 404
    
    # Crear Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Árboles"
    
    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2d6a4f", end_color="2d6a4f", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Escribir encabezados
    columnas = list(datos[0].keys())
    for col_num, columna in enumerate(columnas, 1):
        cell = ws.cell(row=1, column=col_num, value=columna)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Escribir datos
    for row_num, fila in enumerate(datos, 2):
        for col_num, columna in enumerate(columnas, 1):
            cell = ws.cell(row=row_num, column=col_num, value=fila[columna])
            cell.border = border
    
    # Ajustar ancho de columnas automáticamente
    for col_num, columna in enumerate(columnas, 1):
        max_length = len(str(columna))
        for fila in datos:
            value_length = len(str(fila[columna] if fila[columna] is not None else ''))
            if value_length > max_length:
                max_length = value_length
        ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = min(max_length + 3, 50)
    
    # Congelar la fila de encabezados
    ws.freeze_panes = "A2"
    
    # Guardar en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"arboles_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

# ============ EXPORTAR ESPECIES ============
@export_bp.route('/especies/csv', methods=['GET'])
def exportar_especies_csv():
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("""
        SELECT EspecieID AS ID, NombreComun, NombreCientifico, Familia, Categoria
        FROM Especies
        ORDER BY NombreComun
    """)
    datos = cursor.fetchall()
    cursor.close()
    conn.close()
    
    if not datos:
        return jsonify({'error': 'No hay datos para exportar'}), 404
    
    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.DictWriter(output, fieldnames=datos[0].keys())
    writer.writeheader()
    writer.writerows(datos)
    
    output.seek(0)
    bytes_output = io.BytesIO(output.getvalue().encode('utf-8'))
    filename = f"especies_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    
    return send_file(bytes_output, mimetype='text/csv', as_attachment=True, download_name=filename)

@export_bp.route('/especies/xlsx', methods=['GET'])
def exportar_especies_xlsx():
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("""
        SELECT EspecieID AS ID, NombreComun, NombreCientifico, Familia, Categoria
        FROM Especies
        ORDER BY NombreComun
    """)
    datos = cursor.fetchall()
    cursor.close()
    conn.close()
    
    if not datos:
        return jsonify({'error': 'No hay datos para exportar'}), 404
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Especies"
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2d6a4f", end_color="2d6a4f", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    columnas = list(datos[0].keys())
    for col_num, columna in enumerate(columnas, 1):
        cell = ws.cell(row=1, column=col_num, value=columna)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    for row_num, fila in enumerate(datos, 2):
        for col_num, columna in enumerate(columnas, 1):
            cell = ws.cell(row=row_num, column=col_num, value=fila[columna])
            cell.border = border
    
    for col_num, columna in enumerate(columnas, 1):
        max_length = len(str(columna))
        for fila in datos:
            value_length = len(str(fila[columna] if fila[columna] is not None else ''))
            if value_length > max_length:
                max_length = value_length
        ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = min(max_length + 3, 50)
    
    ws.freeze_panes = "A2"
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"especies_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )